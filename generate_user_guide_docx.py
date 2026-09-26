import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def generate_docx():
    doc = docx.Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # Color definitions
    HEX_PRIMARY = "0F172A"
    HEX_SECONDARY = "0284C7"
    HEX_LIGHT = "F8FAFC"
    
    # Title
    title = doc.add_paragraph()
    title_run = title.add_run("SIRAJ UPS & SOLAR")
    title_run.font.name = "Arial"
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42)
    
    subtitle = doc.add_paragraph()
    sub_run = subtitle.add_run("Official User Manual & Complete Operating Guide v1.0")
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(12)
    sub_run.font.bold = True
    sub_run.font.color.rgb = RGBColor(2, 132, 199)

    doc.add_paragraph("=" * 65)

    # Section 1
    h1 = doc.add_heading(level=1)
    r1 = h1.add_run("1. System Overview & Access Details")
    r1.font.color.rgb = RGBColor(15, 23, 42)

    p1 = doc.add_paragraph(
        "Welcome to the official user manual for Siraj UPS and Solar software. "
        "This system manages counter sales, POS billing, inventory stock, customer credit (udhaar), "
        "purchases, expenses, repair services, and financial profit reporting."
    )

    # Credentials Table
    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    
    data = [
        ("Access Method", "Details / Location"),
        ("Desktop Shortcut", "Double-click 'Siraj UPS and Solar' icon on Windows Desktop"),
        ("Standalone EXE App", "SirajUPS_Software.exe inside 'Siraj UPS Software Portable' folder"),
        ("Admin Username", "admin"),
        ("Admin Password", "admin1234"),
        ("Local Web URL", "http://127.0.0.1:8000 (Main Shop Computer)")
    ]

    for i, (col1, col2) in enumerate(data):
        row = table.rows[i]
        row.cells[0].text = col1
        row.cells[1].text = col2
        if i == 0:
            set_cell_background(row.cells[0], HEX_PRIMARY)
            set_cell_background(row.cells[1], HEX_PRIMARY)
            row.cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            row.cells[0].paragraphs[0].runs[0].font.bold = True
            row.cells[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            row.cells[1].paragraphs[0].runs[0].font.bold = True

    doc.add_paragraph()

    # Section 2
    h2 = doc.add_heading(level=1)
    r2 = h2.add_run("2. Core Modules Guide")
    r2.font.color.rgb = RGBColor(15, 23, 42)

    # Modules
    modules = [
        ("A. POS Counter Billing", [
            "Select items by name, SKU, or search bar.",
            "Adjust quantities, unit prices, or apply discounts.",
            "Select Walk-in customer or a registered Udhaar customer.",
            "Choose Cash, Online, or Credit payment method.",
            "Click 'Complete Sale' to generate and print PDF invoice with shop logo."
        ]),
        ("B. Products & Stock Inventory", [
            "Add/edit products with SKU, purchase cost, selling price, and min alert level.",
            "Use '+ New Category' modal to quickly add custom product categories.",
            "Click stock badges for quick stock additions/reductions with notes."
        ]),
        ("C. Customer & Udhaar Ledger", [
            "View customer contact information and real-time total credit balance.",
            "Click 'Record Payment' when a customer returns cash to clear balance.",
            "Download or print complete customer account statement."
        ]),
        ("D. Suppliers & Purchases", [
            "Add battery & solar suppliers/distributors.",
            "Log purchases to automatically increase inventory stock levels.",
            "Track outstanding payments owed to suppliers."
        ]),
        ("E. Repair Services & Expenses", [
            "Log shop expenses (rent, electric bills, staff salaries, tea/food).",
            "Track UPS/Inverter repair jobs, customer equipment, repair status, and charges."
        ])
    ]

    for title, points in modules:
        sh = doc.add_heading(level=2)
        sr = sh.add_run(title)
        sr.font.color.rgb = RGBColor(2, 132, 199)
        for pt in points:
            doc.add_paragraph(pt, style='List Bullet')

    # Section 3
    h3 = doc.add_heading(level=1)
    r3 = h3.add_run("3. Multi-PC & Laptop Usage")
    r3.font.color.rgb = RGBColor(15, 23, 42)

    table2 = doc.add_table(rows=3, cols=2)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    table2.style = 'Table Grid'
    
    data2 = [
        ("Method", "Instructions"),
        ("Multiple Laptops/Mobiles on Shop Wi-Fi", "Find Main PC IP (e.g. 192.168.1.10). Open browser on other devices and go to http://192.168.1.10:8000 while software is running."),
        ("Copy to another Standalone PC (USB)", "Copy 'Siraj UPS Software Portable' folder to USB Drive. Paste on any Windows PC and double click 'SirajUPS_Software.exe'.")
    ]

    for i, (c1, c2) in enumerate(data2):
        row = table2.rows[i]
        row.cells[0].text = c1
        row.cells[1].text = c2
        if i == 0:
            set_cell_background(row.cells[0], HEX_SECONDARY)
            set_cell_background(row.cells[1], HEX_SECONDARY)
            row.cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            row.cells[0].paragraphs[0].runs[0].font.bold = True
            row.cells[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            row.cells[1].paragraphs[0].runs[0].font.bold = True

    output_docx = r"C:\Users\siraj\Desktop\Siraj_UPS_Solar_User_Guide.docx"
    doc.save(output_docx)
    print(f"Word User Guide DOCX generated successfully at: {output_docx}")

if __name__ == "__main__":
    generate_docx()
